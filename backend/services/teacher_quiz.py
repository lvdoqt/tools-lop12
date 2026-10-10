import io
import json
import os

import requests
from fastapi import APIRouter, Depends, File, Header, HTTPException, UploadFile
from pydantic import BaseModel, Field

try:
    from .omr_engine import process_omr_sheet
except ImportError:
    from services.omr_engine import process_omr_sheet

router = APIRouter(prefix='/api/teacher')


def config():
    url = os.getenv('SUPABASE_URL', '').strip().rstrip('/')
    secret = (os.getenv('SUPABASE_SECRET_KEY', '').strip() or
              os.getenv('SUPABASE_SERVICE_ROLE_KEY', '').strip())
    public = os.getenv('SUPABASE_ANON_KEY', '').strip()
    if not url or not secret or not public:
        raise HTTPException(503, 'Cần cấu hình SUPABASE_URL, SUPABASE_ANON_KEY và SUPABASE_SECRET_KEY ở backend.')
    return url, secret, public


def auth_request(path, body, public):
    url, _, _ = config()
    response = requests.post(f'{url}/auth/v1/{path}', json=body,
        headers={'apikey': public, 'Content-Type': 'application/json'}, timeout=20)
    try:
        payload = response.json()
    except ValueError:
        payload = {}
    if not response.ok:
        raise HTTPException(response.status_code, payload.get('msg') or payload.get('message') or payload.get('error_description') or 'Không thể xác thực tài khoản.')
    return payload


def db(method, table, params=None, body=None):
    url, key, _ = config()
    headers = {'apikey': key, 'Prefer': 'return=representation'}
    if not key.startswith('sb_secret_'):
        headers['Authorization'] = f'Bearer {key}'
    response = requests.request(method, f'{url}/rest/v1/{table}', params=params,
        json=body, headers=headers, timeout=20)
    try:
        payload = response.json() if response.content else []
    except ValueError:
        payload = []
    if not response.ok:
        raise HTTPException(502, 'Supabase từ chối thao tác. Kiểm tra SQL thiết lập công cụ chấm trắc nghiệm.')
    return payload


def current_teacher(authorization: str = Header(default='')):
    url, secret, public = config()
    if not authorization.lower().startswith('bearer '):
        raise HTTPException(401, 'Vui lòng đăng nhập tài khoản giáo viên.')
    token = authorization[7:].strip()
    response = requests.get(f'{url}/auth/v1/user', headers={'apikey': public,
        'Authorization': f'Bearer {token}'}, timeout=15)
    if not response.ok:
        raise HTTPException(401, 'Phiên đăng nhập hết hạn. Vui lòng đăng nhập lại.')
    user = response.json()
    teacher_id = user.get('id')
    rows = db('GET', 'teacher_profiles', {'user_id': f'eq.{teacher_id}', 'select': 'user_id'})
    if not rows:
        raise HTTPException(403, 'Tài khoản này chưa được đăng ký vai trò giáo viên.')
    return teacher_id


class Signup(BaseModel):
    email: str = Field(min_length=5, max_length=254)
    password: str = Field(min_length=8, max_length=128)
    full_name: str = Field(min_length=2, max_length=120)


class Login(BaseModel):
    email: str
    password: str


class Exam(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    answers: list[str] = Field(min_length=1, max_length=120)


class Scan(BaseModel):
    image: str = Field(min_length=100, max_length=8_000_000)


@router.post('/signup')
def signup(data: Signup):
    _, _, public = config()
    result = auth_request('signup', {'email': data.email.strip().lower(), 'password': data.password,
        'data': {'full_name': data.full_name.strip()}}, public)
    user = result.get('user') or result
    if user.get('id'):
        db('POST', 'teacher_profiles', body={'user_id': user['id'], 'full_name': data.full_name.strip()})
    return {'message': 'Đã tạo tài khoản giáo viên. Nếu Supabase yêu cầu xác minh email, hãy xác minh trước khi đăng nhập.',
            'session': result.get('session')}


@router.post('/login')
def login(data: Login):
    _, _, public = config()
    result = auth_request('token?grant_type=password', {'email': data.email.strip().lower(), 'password': data.password}, public)
    if not result.get('access_token'):
        raise HTTPException(401, 'Hãy xác minh email trước khi đăng nhập.')
    return result


@router.get('/exams')
def exams(teacher_id: str = Depends(current_teacher)):
    return db('GET', 'teacher_exams', {'teacher_id': f'eq.{teacher_id}', 'select': '*', 'order': 'created_at.desc'})


def parse_answer_file(content: bytes, filename: str) -> list[dict]:
    filename_lower = filename.lower()
    variants = []
    base_title = filename.rsplit('.', 1)[0]

    # 1. EXCEL FILES (.xlsx, .xlsm, .xltx)
    if filename_lower.endswith(('.xlsx', '.xlsm', '.xltx')):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb['Dữ liệu'] if 'Dữ liệu' in wb.sheetnames else wb.active

        rows = []
        for r in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != '' for cell in r):
                rows.append([str(c).strip() if c is not None else '' for c in r])

        if not rows:
            raise ValueError('File Excel không có dữ liệu.')

        header = rows[0]
        # TNMaker matrix format: Row 0 has 'Câu\Mã đề', columns 1..N are exam codes
        if len(header) >= 2 and any(k in header[0].lower() for k in ('câu', 'cau', 'mã', 'ma', 'stt')):
            for col_idx in range(1, len(header)):
                code = header[col_idx]
                if not code:
                    continue
                answers = []
                for row in rows[1:]:
                    if col_idx < len(row):
                        val = row[col_idx]
                        if val:
                            answers.append(val.upper())
                if answers:
                    variants.append({
                        'code': code,
                        'title': f'{base_title} - Mã {code}',
                        'answers': answers,
                        'total': len(answers)
                    })
        else:
            # Single-exam 2-column format (Col A: Câu, Col B: Đáp án)
            ans_list = []
            for idx, r in enumerate(rows):
                val = r[1] if len(r) > 1 else r[0]
                if idx == 0 and any(k in val.lower() for k in ('đáp án', 'dap an', 'answer', 'phương án')):
                    continue
                if val:
                    ans_list.append(val.upper())
            if ans_list:
                variants.append({
                    'code': '0101',
                    'title': base_title,
                    'answers': ans_list,
                    'total': len(ans_list)
                })

    # 2. JSON FILES (.json)
    elif filename_lower.endswith('.json'):
        try:
            data = json.loads(content.decode('utf-8'))
        except Exception as exc:
            raise ValueError(f'File JSON không hợp lệ: {exc}') from exc

        # Format A: Doi_chieu_cau_goc.json from word_shuffle
        # {"0101": {"I": [{"cau_moi": 1, "dap_an": "A"}, ...], "II": [...]}}
        if isinstance(data, dict):
            is_mapping = False
            for k, v in data.items():
                if isinstance(v, dict) and any(sec in v for sec in ('I', 'II', 'III', '1', '2')):
                    is_mapping = True
                    answers = []
                    for sec in ('I', 'II', 'III', '1', '2'):
                        if sec in v and isinstance(v[sec], list):
                            for q in v[sec]:
                                if isinstance(q, dict) and 'dap_an' in q:
                                    answers.append(str(q['dap_an']).strip().upper())
                                elif isinstance(q, str):
                                    answers.append(q.strip().upper())
                    if answers:
                        variants.append({
                            'code': str(k),
                            'title': f'{base_title} - Mã {k}',
                            'answers': answers,
                            'total': len(answers)
                        })
            if not is_mapping:
                # Format B: {"0101": ["A", "B", ...], "0102": ["B", "C", ...]}
                all_code_lists = all(isinstance(v, list) for v in data.values())
                if all_code_lists and len(data) > 0 and 'answers' not in data:
                    for k, ans_list in data.items():
                        cleaned = [str(x).strip().upper() for x in ans_list if str(x).strip()]
                        if cleaned:
                            variants.append({
                                'code': str(k),
                                'title': f'{base_title} - Mã {k}',
                                'answers': cleaned,
                                'total': len(cleaned)
                            })
                else:
                    # Format C: {"title": "...", "answers": ["A", "B", ...]}
                    if 'answers' in data and isinstance(data['answers'], list):
                        cleaned = [str(x).strip().upper() for x in data['answers'] if str(x).strip()]
                        variants.append({
                            'code': str(data.get('code') or '0101'),
                            'title': str(data.get('title') or base_title),
                            'answers': cleaned,
                            'total': len(cleaned)
                        })
        elif isinstance(data, list):
            cleaned = [str(x).strip().upper() for x in data if str(x).strip()]
            variants.append({
                'code': '0101',
                'title': base_title,
                'answers': cleaned,
                'total': len(cleaned)
            })
    else:
        raise ValueError('Chỉ hỗ trợ file Excel (.xlsx) hoặc JSON (.json).')

    if not variants:
        raise ValueError('Không tìm thấy dữ liệu đáp án hợp lệ trong file đã chọn.')

    return variants


@router.post('/import-file')
async def import_file(file: UploadFile = File(...), teacher_id: str = Depends(current_teacher)):
    content = await file.read()
    try:
        variants = parse_answer_file(content, file.filename)
    except Exception as exc:
        raise HTTPException(422, str(exc)) from exc
    return {
        'filename': file.filename,
        'variants': variants
    }


class BatchExams(BaseModel):
    exams: list[Exam]


@router.post('/exams/batch')
def create_exams_batch(batch: BatchExams, teacher_id: str = Depends(current_teacher)):
    created = []
    for item in batch.exams:
        if not item.answers or any(not str(a).strip() for a in item.answers):
            continue
        rows = db('POST', 'teacher_exams', body={
            'teacher_id': teacher_id,
            'title': item.title.strip(),
            'answers': [str(a).strip().upper() for a in item.answers]
        })
        if rows:
            created.append(rows[0])
    return created


@router.post('/exams')
def create_exam(exam: Exam, teacher_id: str = Depends(current_teacher)):
    if any(not str(answer).strip() for answer in exam.answers):
        raise HTTPException(422, 'Mỗi câu cần có đáp án hợp lệ.')
    rows = db('POST', 'teacher_exams', body={'teacher_id': teacher_id, 'title': exam.title.strip(),
        'answers': [str(a).strip().upper() for a in exam.answers]})
    return rows[0]



@router.post('/exams/{exam_id}/scan')
def scan_exam(exam_id: str, scan: Scan, teacher_id: str = Depends(current_teacher)):
    rows = db('GET', 'teacher_exams', {'id': f'eq.{exam_id}', 'teacher_id': f'eq.{teacher_id}', 'select': '*'})
    if not rows:
        raise HTTPException(404, 'Không tìm thấy đề thi.')
    answers = rows[0]['answers']
    try:
        omr_result = process_omr_sheet(scan.image, answers)
    except ValueError as val_err:
        raise HTTPException(422, str(val_err))
    except Exception as exc:
        raise HTTPException(422, f'Không nhận diện được phiếu thi: {str(exc)}') from exc

    result = {
        'exam_id': exam_id,
        'answers': omr_result['answers'],
        'score': omr_result['score'],
        'total': omr_result['total'],
        'uncertain_questions': omr_result['uncertain_questions'],
        'annotated_image': omr_result['annotated_image'],
    }
    saved = db('POST', 'teacher_results', body={
        'teacher_id': teacher_id,
        'exam_id': exam_id,
        'answers': omr_result['answers'],
        'score': omr_result['score'],
        'total': omr_result['total'],
        'uncertain_questions': omr_result['uncertain_questions'],
    })
    result['id'] = saved[0]['id'] if saved else None
    return result


class UpdateResult(BaseModel):
    answers: list[str]


@router.patch('/results/{result_id}')
def update_result(result_id: str, data: UpdateResult, teacher_id: str = Depends(current_teacher)):
    results = db('GET', 'teacher_results', {'id': f'eq.{result_id}', 'teacher_id': f'eq.{teacher_id}', 'select': '*'})
    if not results:
        raise HTTPException(404, 'Không tìm thấy kết quả chấm.')
    current_res = results[0]
    exams = db('GET', 'teacher_exams', {'id': f'eq.{current_res["exam_id"]}', 'select': '*'})
    if not exams:
        raise HTTPException(404, 'Không tìm thấy đề thi tương ứng.')
    expected = exams[0]['answers']
    new_answers = [a.upper().strip() for a in data.answers]
    new_score = sum(got == exp for got, exp in zip(new_answers, expected) if got != '')

    db('PATCH', 'teacher_results', {'id': f'eq.{result_id}', 'teacher_id': f'eq.{teacher_id}'}, body={
        'answers': new_answers,
        'score': new_score,
        'uncertain_questions': [],
    })
    return {
        'id': result_id,
        'score': new_score,
        'total': len(expected),
        'answers': new_answers,
        'uncertain_questions': [],
    }


@router.delete('/results/{result_id}')
def delete_result(result_id: str, teacher_id: str = Depends(current_teacher)):
    db('DELETE', 'teacher_results', {'id': f'eq.{result_id}', 'teacher_id': f'eq.{teacher_id}'})
    return {'ok': True}


@router.delete('/exams/{exam_id}')
def delete_exam(exam_id: str, teacher_id: str = Depends(current_teacher)):
    db('DELETE', 'teacher_exams', {'id': f'eq.{exam_id}', 'teacher_id': f'eq.{teacher_id}'})
    return {'ok': True}



@router.get('/results')
def results(teacher_id: str = Depends(current_teacher)):
    return db('GET', 'teacher_results', {'teacher_id': f'eq.{teacher_id}',
        'select': 'id,exam_id,score,total,created_at,uncertain_questions', 'order': 'created_at.desc', 'limit': 100})

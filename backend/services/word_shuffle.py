"""Shuffle intact OOXML question blocks, retaining MathType, equations and images."""
from copy import deepcopy
from dataclasses import dataclass, field, replace
from io import BytesIO
import json
import math
import random
import re
import zipfile

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W, 'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
SECTION = re.compile(r'^\s*PHẦN\s+(III|II|I)\b', re.I)
QUESTION = re.compile(r'^\s*Câu\s+(\d+)\s*[.:]', re.I)
SOLUTION = re.compile(r'^\s*Lời\s+giải\s*[:.]?\s*$', re.I)
SHORT = re.compile(r'^\s*(?:Đáp\s*án|Đáp\s*số|Trả\s*lời)\s*:\s*([+\-−]?\d+(?:[.,]\d+)?)\s*[.]?\s*$', re.I)
TF_SOLUTION = re.compile(r'(?<!\w)([abcd])\s*\)\s*[:.\-]?\s*(ĐÚNG|SAI)\b', re.I)
MCQ_LABEL = re.compile(r'(?<!\S)([ABCD])\s*[.)]')
MAX_UPLOAD = 4 * 1024 * 1024
MAX_OUTPUT = 4 * 1024 * 1024


def tag(name):
    return f'{{{W}}}{name}'


def text_map(element):
    """Visible text and owning runs; labels may span several Word runs."""
    text, owners = [], []
    for node in element.iter():
        if node.tag in (tag('t'), f'{{{NS["m"]}}}t'):
            value = node.text or ''
        elif node.tag in (tag('tab'), tag('br'), tag('cr')):
            if node.tag == tag('tab') and node.getparent().tag == tag('tabs'):
                continue
            value = '\t' if node.tag == tag('tab') else '\n'
        else:
            continue
        parent = node.getparent()
        while parent is not None and parent.tag != tag('r'):
            parent = parent.getparent()
        text.append(value)
        owners.extend([parent] * len(value))
    return ''.join(text), owners


def paragraphs(blocks):
    for block in blocks:
        if block.tag == tag('p'):
            yield block
        else:
            yield from block.iter(tag('p'))


def plain(blocks):
    return '\n'.join(text_map(p)[0] for p in paragraphs(blocks))


@dataclass
class Question:
    number: int
    blocks: list = field(default_factory=list)
    solution: list = field(default_factory=list)
    answer: str = ''
    options: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    option_order: tuple = (0, 1, 2, 3)


class Exam:
    def __init__(self, content):
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 6000 or sum(i.file_size for i in entries) > 64 * 1024 * 1024:
                    raise ValueError('Tài liệu Word quá lớn sau giải nén (tối đa 64 MB).')
                if len({i.filename for i in entries}) != len(entries):
                    raise ValueError('File Word có các thành phần trùng tên.')
                self.files = {i.filename: archive.read(i) for i in entries}
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            self.root = etree.fromstring(self.files['word/document.xml'], parser)
            if self.root.getroottree().docinfo.doctype:
                raise ValueError('XML của tài liệu Word không hợp lệ.')
            self.styles = {}
            if 'word/styles.xml' in self.files:
                styles = etree.fromstring(self.files['word/styles.xml'], parser)
                self.styles = {s.get(tag('styleId')): s for s in styles.findall('w:style', NS)}
        except (zipfile.BadZipFile, KeyError, etree.XMLSyntaxError, RuntimeError) as exc:
            raise ValueError('Không đọc được DOCX. Hãy mở bằng Word và lưu lại dạng .docx.') from exc
        self.body = self.root.find('w:body', NS)
        if self.body is None:
            raise ValueError('File Word không có nội dung.')
        if self.body.xpath('.//w:ins | .//w:del | .//w:sdt | .//w:altChunk', namespaces=NS):
            raise ValueError('Hãy chấp nhận các thay đổi và chuyển content control thành văn bản trước khi trộn.')
        self.prefix, self.headings = [], {s: [] for s in ('I', 'II', 'III')}
        self.sections = {s: [] for s in self.headings}
        self.errors, self.warnings = [], []
        section, current, solving = None, None, False
        seen = []
        self.sectpr = None
        for block in self.body:
            if block.tag == tag('sectPr'):
                self.sectpr = block
                continue
            if block.tag not in (tag('p'), tag('tbl'), tag('bookmarkStart'), tag('bookmarkEnd')):
                raise ValueError('Tài liệu có khối nội dung chưa hỗ trợ. Hãy lưu câu hỏi bằng đoạn văn Word.')
            value = text_map(block)[0] if block.tag == tag('p') else ''
            heading, question = SECTION.match(value), QUESTION.match(value)
            if heading:
                section = heading[1].upper()
                if section in seen:
                    raise ValueError(f'Phần {section} xuất hiện nhiều lần. Mỗi file chỉ chứa một đề gốc.')
                seen.append(section)
                current, solving = None, False
                self.headings[section].append(block)
            elif question:
                if section is None:
                    raise ValueError('Có câu hỏi trước tiêu đề PHẦN I.')
                current = Question(int(question[1]), [block])
                self.sections[section].append(current)
                solving = False
            elif current is not None:
                if SOLUTION.match(value):
                    solving = True
                (current.solution if solving else current.blocks).append(block)
            elif section:
                self.headings[section].append(block)
            else:
                self.prefix.append(block)
            if block.tag == tag('tbl') and any(QUESTION.match(text_map(p)[0]) for p in paragraphs([block])):
                raise ValueError('Câu hỏi nằm trong bảng chưa được hỗ trợ. Hãy đưa nhãn Câu n ra ngoài bảng; bảng trong nội dung vẫn được giữ.')
        if seen != ['I', 'II', 'III']:
            raise ValueError('Cần đủ ba tiêu đề PHẦN I, PHẦN II, PHẦN III theo đúng thứ tự.')
        for section, questions in self.sections.items():
            if not questions:
                self.errors.append(f'Phần {section} chưa có câu hỏi.')
            numbers = [q.number for q in questions]
            if len(set(numbers)) != len(numbers):
                self.errors.append(f'Phần {section} có số câu trùng nhau.')
            for question in questions:
                self.read_answer(section, question)
                self.errors.extend(f'Phần {section}, câu {question.number}: {e}' for e in question.errors)
                has_equation = any(b.xpath('.//w:object | .//m:oMath | .//w:drawing | .//w:pict', namespaces=NS) for b in question.solution)
                if not has_equation and (not plain(question.solution).strip() or SOLUTION.fullmatch(plain(question.solution).strip())):
                    self.warnings.append(f'Phần {section}, câu {question.number}: chưa có nội dung lời giải.')
        if sum(map(len, self.sections.values())) > 200:
            raise ValueError('Mỗi đề tối đa 200 câu.')

    def underlined(self, run):
        if run is None:
            return False
        direct = run.find('w:rPr/w:u', NS)
        if direct is not None:
            return direct.get(tag('val'), 'single') not in ('none', '0', 'false')
        ids = []
        style = run.find('w:rPr/w:rStyle', NS)
        if style is not None:
            ids.append(style.get(tag('val')))
        paragraph = run.getparent()
        while paragraph is not None and paragraph.tag != tag('p'):
            paragraph = paragraph.getparent()
        if paragraph is not None:
            style = paragraph.find('w:pPr/w:pStyle', NS)
            if style is not None:
                ids.append(style.get(tag('val')))
        for style_id in ids:
            visited = set()
            while style_id in self.styles and style_id not in visited:
                visited.add(style_id)
                style = self.styles[style_id]
                underline = style.find('w:rPr/w:u', NS)
                if underline is not None:
                    return underline.get(tag('val'), 'single') not in ('none', '0', 'false')
                base = style.find('w:basedOn', NS)
                style_id = base.get(tag('val')) if base is not None else None
        return False

    def read_answer(self, section, question):
        if section == 'III':
            matches = [SHORT.fullmatch(text_map(p)[0]) for p in paragraphs(question.solution)]
            answers = [m[1].replace('−', '-').replace('.', ',').lstrip('+') for m in matches if m]
            if len(answers) != 1:
                question.errors.append('Cần đúng một dòng “Đáp án: số”, “Đáp số: số” hoặc “Trả lời: số” dưới Lời giải.')
            else:
                question.answer = answers[0]
            return
        pattern = re.compile(r'(?<!\S)([ABCD])\s*[.)]' if section == 'I' else r'(?<!\S)([abcd])\s*\)')
        for paragraph in paragraphs(question.blocks):
            value, owners = text_map(paragraph)
            matches = list(pattern.finditer(value))
            for i, match in enumerate(matches):
                question.options.append({
                    'label': match[1], 'correct': self.underlined(owners[match.start(1)]),
                    'text': value[match.end():matches[i + 1].start() if i + 1 < len(matches) else len(value)].strip(),
                })
        expected = list('ABCD' if section == 'I' else 'abcd')
        if [o['label'] for o in question.options] != expected:
            question.errors.append('Không tách được đủ bốn nhãn ' + ', '.join(expected) + ' theo thứ tự.')
        elif section == 'I':
            correct = [o['label'] for o in question.options if o['correct']]
            if len(correct) != 1:
                question.errors.append('Cần gạch chân đúng một nhãn A/B/C/D của phương án đúng.')
            else:
                question.answer = correct[0]
        else:
            explicit = {}
            for match in TF_SOLUTION.finditer(plain(question.solution)):
                label, correct = match[1].lower(), match[2].upper() == 'ĐÚNG'
                if label in explicit and explicit[label] != correct:
                    question.errors.append(f'Lời giải ghi cả ĐÚNG và SAI cho ý {label}).')
                explicit[label] = correct
            has_underlines = any(o['correct'] for o in question.options)
            if explicit and not has_underlines and len(explicit) != 4:
                missing = ', '.join(label + ')' for label in 'abcd' if label not in explicit)
                question.errors.append(f'Lời giải chưa ghi ĐÚNG/SAI cho ý {missing}; đề không có nhãn gạch chân để đối chiếu.')
            for option in question.options:
                label = option['label']
                if label in explicit:
                    if option['correct'] and not explicit[label]:
                        question.errors.append(f'Ý {label}) được gạch chân nhưng Lời giải ghi SAI.')
                    option['correct'] = explicit[label]
            if question.errors:
                return
            question.answer = ''.join('Đ' if o['correct'] else 'S' for o in question.options)

    def summary(self):
        return {'total': sum(map(len, self.sections.values())), 'errors': self.errors, 'warnings': self.warnings,
                'sections': [{'part': part, 'count': len(questions), 'questions': [
                    {'number': q.number, 'text': plain(q.blocks), 'answer': q.answer,
                     'options': q.options, 'solution': plain(q.solution), 'errors': q.errors}
                    for q in questions]} for part, questions in self.sections.items()]}


def renumber(block, number):
    nodes = list(block.iter(tag('t')))
    value = ''.join(n.text or '' for n in nodes)
    match = QUESTION.match(value)
    if not match:
        raise ValueError('Không thể đánh lại số câu hỏi.')
    start, end = match.span(1)
    offset = 0
    for node in nodes:
        old = node.text or ''
        left, right = max(0, start - offset), min(len(old), end - offset)
        if left < right:
            node.text = old[:left] + (str(number) if offset <= start < offset + len(old) else '') + old[right:]
        offset += len(old)


def rich_slice(p, start, end):
    """Slice a paragraph by visible text offsets without converting its math or OLE."""
    total = len(text_map(p)[0])
    position = 0

    def visit(node):
        nonlocal position
        if node.tag in (tag('pPr'), tag('rPr')):
            return deepcopy(node)
        length = len(text_map(node)[0])
        if not length:
            return deepcopy(node) if start <= position < end or position == end == total else None
        if node.tag in (tag('t'), f'{{{NS["m"]}}}t', tag('tab'), tag('br'), tag('cr')):
            left, right = max(0, start - position), min(length, end - position)
            position += length
            if left >= right:
                return None
            result = deepcopy(node)
            if node.tag in (tag('t'), f'{{{NS["m"]}}}t'):
                result.text = (node.text or '')[left:right]
                result.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            return result
        result = deepcopy(node)
        for child in list(result):
            result.remove(child)
        for child in node:
            copied = visit(child)
            if copied is not None:
                result.append(copied)
        return result if any(c.tag not in (tag('pPr'), tag('rPr')) for c in result) else None

    result = visit(p)
    if result is not None:
        for attribute in list(result.attrib):
            if attribute.endswith('}paraId') or attribute.endswith('}textId'):
                del result.attrib[attribute]
    return result


def trim_option(p):
    # Separators belong to layout slots, not to the answer being moved.
    for node in reversed(list(p.iter())):
        if node.tag in (tag('bookmarkStart'), tag('bookmarkEnd'), tag('r'), tag('rPr')):
            continue
        if node.tag == tag('t'):
            node.text = (node.text or '').rstrip()
            if node.text:
                break
        elif node.tag == tag('tab'):
            node.getparent().remove(node)
        elif node.tag not in (tag('p'),) and not node.tag.startswith(f'{{{W}}}'):
            break
        elif node.tag in (tag('object'), tag('drawing'), tag('pict')):
            break


def mark_tf_answers(blocks, answer):
    """Mark resolved answers even when the source only gives them in its solution."""
    for p in paragraphs(blocks):
        value, _ = text_map(p)
        matches = list(re.finditer(r'(?<!\S)([abcd])\s*\)', value))
        if not matches:
            continue
        pieces, cursor = [], 0
        for match in matches:
            start, end = match.span(1)
            if cursor < start:
                pieces.append(rich_slice(p, cursor, start))
            label = rich_slice(p, start, end)
            for run in label.iter(tag('r')):
                if not any(n.text for n in run.iter(tag('t'))):
                    continue
                props = run.find('w:rPr', NS)
                if props is None:
                    props = etree.Element(tag('rPr'))
                    run.insert(0, props)
                underline = props.find('w:u', NS)
                if underline is None:
                    underline = etree.SubElement(props, tag('u'))
                underline.set(tag('val'), 'single' if answer['abcd'.index(match[1])] == 'Đ' else 'none')
            pieces.append(label)
            cursor = end
        if cursor < len(value):
            pieces.append(rich_slice(p, cursor, len(value)))
        for child in list(p):
            if child.tag != tag('pPr'):
                p.remove(child)
        for piece in pieces:
            if piece is not None:
                p.extend(child for child in piece if child.tag != tag('pPr'))
    return blocks


def shuffled_blocks(question):
    """Keep the source's 4/2/1-option rows; move multi-paragraph answers together."""
    stem, choices, groups = [], [], []
    simple = True

    def blocks_with_options(blocks):
        for block in blocks:
            if block.tag == tag('tbl') and any(MCQ_LABEL.search(text_map(p)[0]) for p in paragraphs([block])):
                # Option tables become paragraph rows, keeping all rich cell content.
                for cell in block.findall('w:tr/w:tc', NS):
                    yield from blocks_with_options([b for b in cell if b.tag != tag('tcPr')])
            else:
                yield block

    for block in blocks_with_options(question.blocks):
        value = text_map(block)[0] if block.tag == tag('p') else ''
        matches = list(MCQ_LABEL.finditer(value))
        if not matches:
            if choices:
                choices[-1].append(deepcopy(block))
                simple = False
            else:
                stem.append(deepcopy(block))
            continue
        prefix = rich_slice(block, 0, matches[0].start())
        if prefix is not None and (text_map(prefix)[0].strip() or prefix.xpath('.//w:object | .//w:drawing | .//m:oMath', namespaces=NS)):
            if choices:
                choices[-1].append(prefix)
                simple = False
            else:
                stem.append(prefix)
        groups.append((block, len(matches)))
        for i, match in enumerate(matches):
            end = matches[i + 1].start() if i + 1 < len(matches) else len(value)
            fragment = rich_slice(block, match.start(), end)
            trim_option(fragment)
            choices.append([fragment])
    if len(choices) != 4:
        raise ValueError(f'Không thể tách phương án Phần I, câu {question.number}.')
    reordered = [choices[index] for index in question.option_order]
    for index, blocks in enumerate(reordered):
        # The label may be a single run or share a run with the whole option.
        node = next(n for n in blocks[0].iter(tag('t')) if n.text)
        node.text = 'ABCD'[index] + node.text[1:]
        run = node.getparent()
        while run is not None and run.tag != tag('r'):
            run = run.getparent()
        if run is not None:
            props = run.find('w:rPr', NS)
            if props is None:
                props = etree.Element(tag('rPr'))
                run.insert(0, props)
            underline = props.find('w:u', NS)
            if underline is None:
                underline = etree.SubElement(props, tag('u'))
            underline.set(tag('val'), 'single' if 'ABCD'[index] == question.answer else 'none')
    if simple:
        offset = 0
        for template, size in groups:
            row = deepcopy(template)
            for child in list(row):
                if child.tag != tag('pPr'):
                    row.remove(child)
            for index in range(size):
                if index:
                    etree.SubElement(etree.SubElement(row, tag('r')), tag('tab'))
                row.extend(child for child in reordered[offset + index][0] if child.tag != tag('pPr'))
            stem.append(row)
            offset += size
    else:
        stem.extend(block for blocks in reordered for block in blocks)
    return stem


def remap_solution(blocks, option_order):
    """Update explicit answer references, without touching mathematical points A–D."""
    mapping = {old: 'ABCD'[new] for new, old in enumerate('ABCD'[i] for i in option_order)}
    pattern = re.compile(r'(?i:\b(?:chọn|đáp\s*án(?:\s*đúng)?(?:\s*là)?|phương\s*án))\s*[:.]?\s*([ABCD])\b')
    for p in paragraphs(blocks):
        nodes = list(p.iter(tag('t')))
        value = ''.join(n.text or '' for n in nodes)
        positions = {m.start(1): mapping[m[1]] for m in pattern.finditer(value)}
        offset = 0
        for node in nodes:
            old = node.text or ''
            node.text = ''.join(positions.get(offset + i, char) for i, char in enumerate(old))
            offset += len(old)
    return blocks


def hide_marks(block, section):
    pattern = re.compile(r'(?<!\S)([ABCD])\s*[.)]' if section == 'I' else r'(?<!\S)([abcd])\s*\)')
    for paragraph in paragraphs([block]):
        value, _ = text_map(paragraph)
        # Clear answer underlining on option runs only; preserve underlined prose in stems.
        matches = list(pattern.finditer(value)) if section != 'III' else []
        if matches:
            for run in paragraph.iter(tag('r')):
                props = run.find('w:rPr', NS)
                if props is None:
                    props = etree.Element(tag('rPr'))
                    run.insert(0, props)
                for underline in list(props.findall('w:u', NS)):
                    props.remove(underline)
                etree.SubElement(props, tag('u')).set(tag('val'), 'none')
        for highlight in list(paragraph.iter(tag('highlight'))):
            highlight.getparent().remove(highlight)


def paragraph(value):
    p = etree.Element(tag('p'))
    run = etree.SubElement(p, tag('r'))
    etree.SubElement(run, tag('t')).text = value
    return p


def heading_paragraph(value, size=26, bold=True, italic=False, color='000000', after=60):
    p = paragraph(value)
    props = etree.Element(tag('pPr'))
    p.insert(0, props)
    etree.SubElement(props, tag('keepNext'))
    etree.SubElement(props, tag('jc')).set(tag('val'), 'center')
    spacing = etree.SubElement(props, tag('spacing'))
    spacing.set(tag('before'), '0')
    spacing.set(tag('after'), str(after))
    spacing.set(tag('line'), '240')
    spacing.set(tag('lineRule'), 'auto')
    run = p.find('w:r', NS)
    rpr = etree.Element(tag('rPr'))
    run.insert(0, rpr)
    fonts = etree.SubElement(rpr, tag('rFonts'))
    for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(tag(name), 'Times New Roman')
    for name, val in (('b', '1' if bold else '0'), ('i', '1' if italic else '0'), ('color', color), ('sz', str(size))):
        etree.SubElement(rpr, tag(name)).set(tag('val'), val)
    return p


def exam_heading(exam, code, department, school):
    width = 10000
    if exam.sectpr is not None:
        page = exam.sectpr.find('w:pgSz', NS)
        margins = exam.sectpr.find('w:pgMar', NS)
        if page is not None and margins is not None:
            width = int(page.get(tag('w'), '11906')) - sum(int(margins.get(tag(side), '850')) for side in ('left', 'right'))
    widths = [int(width * .36), width - int(width * .36)]
    table = etree.Element(tag('tbl'))
    props = etree.SubElement(table, tag('tblPr'))
    table_width = etree.SubElement(props, tag('tblW'))
    table_width.set(tag('w'), str(width))
    table_width.set(tag('type'), 'dxa')
    borders = etree.SubElement(props, tag('tblBorders'))
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        etree.SubElement(borders, tag(side)).set(tag('val'), 'nil')
    etree.SubElement(props, tag('tblLayout')).set(tag('type'), 'fixed')
    grid = etree.SubElement(table, tag('tblGrid'))
    for column_width in widths:
        etree.SubElement(grid, tag('gridCol')).set(tag('w'), str(column_width))
    row = etree.SubElement(table, tag('tr'))
    etree.SubElement(etree.SubElement(row, tag('trPr')), tag('cantSplit'))
    lines = [
        [heading_paragraph(department.upper(), 24, False), heading_paragraph(school.upper(), 24)],
        [heading_paragraph('ĐỀ THI THỬ TN THPT 2027', 28), heading_paragraph('MÔN: TOÁN', 26),
         heading_paragraph('Thời gian làm bài: 90 phút', 24, False, True)],
    ]
    for column_width, content in zip(widths, lines):
        cell = etree.SubElement(row, tag('tc'))
        tcpr = etree.SubElement(cell, tag('tcPr'))
        cell_width = etree.SubElement(tcpr, tag('tcW'))
        cell_width.set(tag('w'), str(column_width))
        cell_width.set(tag('type'), 'dxa')
        etree.SubElement(tcpr, tag('vAlign')).set(tag('val'), 'top')
        cell.extend(content)
    return [table, heading_paragraph(f'Mã đề: {code}', 36, color='FF0000', after=180)]


def render_document(exam, order, code, teacher, department, school):
    root = deepcopy(exam.root)
    body = root.find('w:body', NS)
    for child in list(body):
        body.remove(child)
    body.extend(exam_heading(exam, code, department, school))
    for part, questions in order.items():
        body.extend(deepcopy(exam.headings[part]))
        for number, q in enumerate(questions, 1):
            blocks = shuffled_blocks(q) if part == 'I' else deepcopy(q.blocks)
            if teacher and part == 'II':
                mark_tf_answers(blocks, q.answer)
            renumber(blocks[0], number)
            for block in blocks:
                if not teacher:
                    hide_marks(block, part)
                body.append(block)
            if teacher:
                solution = deepcopy(q.solution)
                body.extend(remap_solution(solution, q.option_order) if part == 'I' else solution)
    if exam.sectpr is not None:
        body.append(deepcopy(exam.sectpr))
    output = BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in exam.files.items():
            archive.writestr(name, etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
                             if name == 'word/document.xml' else data)
    return output.getvalue()


def answer_workbook(variants):
    """TNMaker's official 2025 direct-import matrix, single data sheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Dữ liệu'
    sheet.append(['Câu\\Mã đề'] + [code for code, _ in variants])
    answers = [[q.answer for questions in order.values() for q in questions] for _, order in variants]
    for index, row in enumerate(zip(*answers), 1):
        sheet.append([index, *row])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='975A16')
    for row in sheet.iter_rows(min_col=2):
        for cell in row:
            cell.number_format = '@'
    sheet.freeze_panes = 'B2'
    sheet.column_dimensions['A'].width = 16
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_shuffle(content, count=4, start_code='0101', include_solutions=True, seed=None, variant_index=None,
                   department='SỞ GDĐT ................................', school='TRƯỜNG THPT ........................'):
    if not 1 <= count <= 50:
        raise ValueError('Số đề phải từ 1 đến 50.')
    if not re.fullmatch(r'\d{3,4}', start_code) or int(start_code) + count - 1 > 9999:
        raise ValueError('Mã đề phải có 3–4 chữ số và mã cuối không vượt quá 9999.')
    if variant_index is not None and (seed is None or not 0 <= variant_index < count):
        raise ValueError('Chỉ số mã đề hoặc phiên trộn không hợp lệ.')
    department, school = department.strip(), school.strip()
    if not department or not school or max(len(department), len(school)) > 120:
        raise ValueError('Tên Sở GDĐT và Trường THPT cần có nội dung, tối đa 120 ký tự mỗi dòng.')
    exam = Exam(content)
    if exam.errors:
        raise ValueError('\n'.join(exam.errors))
    possibilities = math.prod(math.factorial(len(q)) for q in exam.sections.values()) * 24 ** len(exam.sections['I'])
    if count > possibilities:
        raise ValueError(f'Đề này chỉ tạo được {possibilities} cách trộn khác nhau.')
    rng = random.Random(seed)
    variants, seen = [], set()
    while len(variants) < count:
        order = {part: rng.sample(q, len(q)) for part, q in exam.sections.items()}
        shuffled = []
        for q in order['I']:
            option_order = tuple(rng.sample(range(4), 4))
            answer = 'ABCD'[option_order.index('ABCD'.index(q.answer))]
            shuffled.append(replace(q, answer=answer, option_order=option_order))
        order['I'] = shuffled
        signature = tuple(tuple((q.number, q.option_order) for q in questions) for questions in order.values())
        if signature in seen:
            continue
        seen.add(signature)
        variants.append((str(int(start_code) + len(variants)).zfill(len(start_code)), order))
    output = BytesIO()
    mapping = {code: {part: [{'cau_moi': i, 'cau_goc': q.number, 'dap_an': q.answer,
                            **({'phuong_an_moi_sang_goc': {new: 'ABCD'[old] for new, old in zip('ABCD', q.option_order)}} if part == 'I' else {})}
                            for i, q in enumerate(questions, 1)] for part, questions in order.items()}
               for code, order in variants}
    selected = variants if variant_index is None else [variants[variant_index]]
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for code, order in selected:
            archive.writestr(f'De_{code}.docx', render_document(exam, order, code, False, department, school))
            if include_solutions:
                archive.writestr(f'loi_giai-{code}.docx', render_document(exam, order, code, True, department, school))
            if output.tell() > MAX_OUTPUT:
                raise ValueError('Một lượt xuất vượt quá 4 MB. Hãy bỏ bản lời giải hoặc giảm dung lượng hình trong đề gốc.')
        archive.writestr('Dap_an_TNMaker_2025.xlsx', answer_workbook(variants))
        archive.writestr('Doi_chieu_cau_goc.json', json.dumps(mapping, ensure_ascii=False, indent=2))
        archive.writestr('Huong_dan.txt', (
            'Excel dùng mẫu nhập trực tiếp TNMaker 2025, một sheet Dữ liệu.\n'
            'Hàng đầu: Câu\\Mã đề và các mã đề. Cột đầu đánh số liên tục qua ba phần.\n'
            'Phần I: A/B/C/D; Phần II: chuỗi bốn ký tự Đ/S theo a,b,c,d; Phần III: số dùng dấu phẩy.\n'
            'Trong TNMaker chọn phiếu 2025, đặt số câu từng phần đúng với đề rồi nhập file Excel.\n'
            'Mẫu chính thức: https://tnmaker.net/nap-dap-an-phieu-tltn-2025/\n'
            'Trộn câu trong cùng phần và trộn A–D ở Phần I; giữ thứ tự a–d ở Phần II.\n'
            'De_*.docx: bản học sinh. loi_giai-*.docx: bản gạch chân đáp án đúng và lời giải.\n'
            + '\n'.join(exam.warnings)).encode('utf-8'))
    if len(output.getvalue()) > MAX_OUTPUT:
        raise ValueError('ZIP vượt quá 4 MB. Hãy giảm số đề hoặc bỏ bản lời giải.')
    return output.getvalue()

"""Validate Quiz Bank JSON and build an exam with two editable equation formats."""

from collections import Counter
from io import BytesIO
import json
import re
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt
from docx_equation.mathtype.ooxml import make_object_run
from lxml import html

from .latex_quiz import MATH_PATTERN
from .word_math import compile_math, equation_element
from .word_images import load_images, svg_png

TYPES = ('mcq', 'msq', 'sa')
SECTIONS = {'mcq': 'PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.', 'msq': 'PHẦN II. Câu trắc nghiệm đúng sai.', 'sa': 'PHẦN III. Câu trắc nghiệm trả lời ngắn.'}
FIELDS = ('question', 'option_a', 'option_b', 'option_c', 'option_d', 'explanation')
MAX_SOURCE = 2 * 1024 * 1024
WORD_MATH_PATTERN = re.compile(
    MATH_PATTERN.pattern + r'|\\begin\{(?P<environment>equation\*?|align\*?|eqnarray\*?|gather\*?|multline\*?)\}.*?\\end\{(?P=environment)\}',
    re.DOTALL,
)


def formula_body(text):
    if text.startswith(r'\begin{'):
        return text
    return text[2:-2] if text.startswith(('$$', r'\(', r'\[')) else text[1:-1]


def content_tree(value, formulas, location):
    def math(match):
        latex = formula_body(match.group())
        if len(latex) > 8000:
            raise ValueError(f'{location}: công thức quá dài.')
        if latex not in formulas:
            if len(formulas) >= 1500:
                raise ValueError('Mỗi lần xuất tối đa 1.500 công thức khác nhau.')
            try:
                formulas[latex] = compile_math(latex)
            except Exception as exc:
                raise ValueError(f'{location}: không chuyển được công thức {latex[:100]!r}: {exc}') from exc
        return f'<quiz-math data-id="{formulas[latex]["id"]}"></quiz-math>'

    protected = WORD_MATH_PATTERN.sub(math, value)
    protected = re.sub(r'\\allowdisplaybreaks(?:\[[^\]]*\])?', '', protected)
    if re.search(r'(?<!\\)\$|\\[\[\]()]', protected):
        raise ValueError(f'{location}: dấu phân cách công thức chưa khép kín.')
    # Only treat actual HTML tags as markup; comparisons such as a < b remain text.
    protected = re.sub(r'<(?!/?[A-Za-z][A-Za-z0-9-]*(?:\s|/?>))', '&lt;', protected)
    root = html.fragment_fromstring(protected, create_parent='div')
    if root.xpath('.//script|.//iframe|.//object|.//embed|.//svg|.//style'):
        raise ValueError(f'{location}: HTML chứa thành phần không hỗ trợ.')
    for image in root.iter('img'):
        src = image.get('src', '').strip()
        if not src or not src.startswith(('https://', 'data:image/')):
            raise ValueError(f'{location}: hình cần src HTTPS hoặc data:image.')
    for table in root.iter('table'):
        rows = table.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
        widths = [len(row.xpath('./td|./th')) for row in rows]
        if not rows or len(rows) > 100 or not widths[0] or max(widths) > 20 or len(set(widths)) != 1:
            raise ValueError(f'{location}: bảng phải có số cột đều nhau (tối đa 20 cột, 100 hàng).')
        if table.xpath('.//table|.//*[@colspan and @colspan!="1"]|.//*[@rowspan and @rowspan!="1"]'):
            raise ValueError(f'{location}: bảng lồng hoặc gộp ô chưa hỗ trợ.')
    return root


def validate_quiz(source, include_answers=False):
    if len(source.encode('utf-8')) > MAX_SOURCE:
        raise ValueError('JSON tối đa 2 MB.')
    try:
        data = json.loads(source.lstrip('\ufeff'))
    except (ValueError, RecursionError) as exc:
        raise ValueError(f'JSON không hợp lệ: {exc}') from exc
    if isinstance(data, dict) and 'quiz' in data:
        data = data['quiz']
    if not isinstance(data, dict) or not isinstance(data.get('questions'), list) or not 1 <= len(data['questions']) <= 100:
        raise ValueError('Cần đối tượng JSON có questions là danh sách từ 1 đến 100 câu.')
    formulas, trees, urls, warnings = {}, [], set(), []
    counts = Counter()
    for index, question in enumerate(data['questions'], 1):
        if not isinstance(question, dict) or question.get('type') not in TYPES:
            raise ValueError(f'Câu {index}: type phải là mcq, msq hoặc sa.')
        kind = question['type']
        counts[kind] += 1
        required = ('question',) + tuple(f'option_{letter}' for letter in 'abcd') if kind != 'sa' else ('question',)
        for key in required:
            if not isinstance(question.get(key), str) or not question[key].strip():
                raise ValueError(f'Câu {index}: thiếu nội dung {key}.')
        answer = question.get('correct_option')
        if answer is not None:
            if not isinstance(answer, str):
                raise ValueError(f'Câu {index}: correct_option phải là chuỗi.')
            if kind == 'mcq' and answer not in 'ABCD' or kind == 'mcq' and len(answer) != 1:
                raise ValueError(f'Câu {index}: đáp án phải là A, B, C hoặc D.')
            if kind == 'msq' and (not re.fullmatch(r'(?:[ABCD](?:,[ABCD])*)?', answer) or len(set(answer.split(','))) != len(answer.split(','))):
                raise ValueError(f'Câu {index}: đáp án đúng/sai phải dạng A,C; để rỗng nếu tất cả sai.')
            if kind == 'sa' and not answer.strip():
                raise ValueError(f'Câu {index}: đáp án trả lời ngắn đang rỗng.')
        elif include_answers:
            raise ValueError(f'Câu {index}: chưa có correct_option để xuất đáp án.')
        active = list(required)
        if include_answers:
            active.append('explanation')
        question_trees = {}
        for key in active:
            value = question.get(key, '')
            if not isinstance(value, str) or len(value) > 100_000:
                raise ValueError(f'Câu {index}: {key} phải là chuỗi tối đa 100.000 ký tự.')
            tree = content_tree(value, formulas, f'Câu {index}, {key}')
            question_trees[key] = tree
            urls.update(node.get('src').strip() for node in tree.iter('img'))
        if include_answers and kind == 'sa':
            question_trees['answer'] = content_tree(answer, formulas, f'Câu {index}, đáp án')
        trees.append(question_trees)
    if len(urls) > 60:
        raise ValueError('Mỗi lần xuất tối đa 60 ảnh khác nhau.')
    standard = all(counts[kind] == count for kind, count in zip(TYPES, (12, 4, 6)))
    if not standard:
        warnings.append('Đề chưa đủ cấu trúc 12 câu chọn đáp án, 4 câu đúng/sai, 6 câu trả lời ngắn. Vẫn xuất toàn bộ câu đã nhập, không tự thêm hoặc bỏ câu.')
    return {'quiz': data, 'trees': trees, 'formulas': formulas, 'urls': urls, 'summary': {
        'title': str(data.get('title', 'Đề thi Toán'))[:200], 'counts': {kind: counts[kind] for kind in TYPES},
        'total': len(trees), 'image_count': len(urls), 'formula_count': len(formulas), 'standard': standard, 'warnings': warnings,
        'formulas': [{key: formula[key] for key in ('id', 'latex', 'mathml')} for formula in formulas.values()],
    }}


def field(paragraph, instruction):
    node = OxmlElement('w:fldSimple')
    node.set(qn('w:instr'), instruction)
    paragraph._p.append(node)


class ExamWriter:
    def __init__(self, prepared, settings, images, previews, mode):
        self.prepared, self.settings, self.images, self.previews, self.mode = prepared, settings, images, previews, mode
        self.formulas = {item['id']: item for item in prepared['formulas'].values()}
        self.doc = Document()
        self.equation_count = 0
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
        section.left_margin, section.right_margin = Cm(2), Cm(1.5)
        section.header_distance, section.footer_distance = Cm(0.6), Cm(0.7)
        normal = self.doc.styles['Normal']
        normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.05
        normal.paragraph_format.widow_control = True
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Trang '); field(footer, 'PAGE'); footer.add_run('/'); field(footer, 'NUMPAGES')
        footer.add_run(f' – Mã đề {settings["code"]}')
        self.doc.core_properties.title = settings['exam_title']
        self.doc.core_properties.subject = 'Đề thi môn Toán'
        # Page fields update during pagination. Preserve embedded equation
        # previews on open instead of activating every installed OLE server.
        update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'false')
        self.doc.settings.element.append(update)

    def paragraph(self, parent=None, text='', bold=False, center=False):
        p = (parent or self.doc).add_paragraph()
        if text:
            p.add_run(text).bold = bold
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def add_equation(self, p, identifier):
        formula = self.formulas[identifier]
        self.equation_count += 1
        if self.mode == 'equation':
            p._p.append(equation_element(formula))
            return
        png, width, height, baseline = self.previews[identifier]
        image_rel, _ = p.part.get_or_add_image(BytesIO(png))
        # Each occurrence owns its storage, so editing the solution cannot
        # change an identical equation in the exam (or another answer choice).
        part = Part(PackURI(f'/word/embeddings/equation-{self.equation_count}.bin'), 'application/vnd.openxmlformats-officedocument.oleObject', formula['ole'], self.doc.part.package)
        ole_rel = self.doc.part.relate_to(part, RT.OLE_OBJECT)
        node = make_object_run(image_rel, ole_rel, width, height, index=self.equation_count, preview_pt_per_px=1, max_width_pt=490)
        if baseline:
            props = OxmlElement('w:rPr'); position = OxmlElement('w:position')
            position.set(qn('w:val'), str(-round(baseline * 2)))
            props.append(position); node.insert(0, props)
        p._p.append(node)

    def content(self, root, parent=None, p=None, bold=False, italic=False):
        parent = parent or self.doc
        def ensure():
            nonlocal p
            if p is None:
                p = self.paragraph(parent)
            return p
        def text(value):
            if not value:
                return
            value = value.replace(r'\%', '%').replace(r'\&', '&').replace(r'\_', '_')
            value = value.replace(r'\lq\lq', '“').replace(r'\rq\rq', '”').replace(r'\lq', '‘').replace(r'\rq', '’')
            value = re.sub(r'\\(?=\s)', '', value)
            for index, line in enumerate(value.split('\n')):
                if index:
                    ensure().add_run().add_break()
                if line:
                    run = ensure().add_run(line)
                    run.bold, run.italic = bold, italic
        text(root.text)
        for node in root:
            tag = node.tag.lower() if isinstance(node.tag, str) else ''
            if tag == 'quiz-math':
                self.add_equation(ensure(), node.get('data-id'))
            elif tag == 'img':
                data, width, height = self.images[node.get('src').strip()]
                scale = min(1, 430 / width, 250 / height)
                picture = self.paragraph(parent, center=True)
                picture.add_run().add_picture(BytesIO(data), width=Pt(width * scale), height=Pt(height * scale))
                p = None
            elif tag == 'table':
                rows = node.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
                table = parent.add_table(rows=len(rows), cols=len(rows[0].xpath('./td|./th')))
                table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for row, source_row in zip(table.rows, rows):
                    no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
                    for cell, source_cell in zip(row.cells, source_row.xpath('./td|./th')):
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        self.content(source_cell, cell, cell.paragraphs[0], bold=source_cell.tag == 'th')
                p = None
            elif tag == 'br':
                ensure().add_run().add_break()
            elif tag in ('p', 'div', 'center', 'li', 'ul', 'ol'):
                p = self.paragraph(parent)
                if tag == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if tag == 'li': p.add_run('• ')
                self.content(node, parent, p, bold, italic)
                p = None
            else:
                p = self.content(node, parent, ensure(), bold or tag in ('b', 'strong'), italic or tag in ('i', 'em'))
            text(node.tail)
        return p

    def option_width(self, root):
        if root.xpath('.//img|.//table|.//br'):
            return 1000
        width = len(''.join(root.itertext())) * 6 + 20
        width += sum(self.previews[node.get('data-id')][1] for node in root.iter('quiz-math'))
        return width

    def options(self, trees, kind):
        width = max(self.option_width(trees[f'option_{letter}']) for letter in 'abcd')
        columns = (4 if width < 110 else 2 if width < 225 else 1) if kind == 'mcq' else 1
        if columns == 1:
            for letter in 'abcd':
                p = self.paragraph(text=f'{letter.upper()}.' if kind == 'mcq' else f'{letter})', bold=True)
                p.add_run(' ')
                p.paragraph_format.left_indent = Cm(0.35)
                p.paragraph_format.space_after = Pt(2)
                self.content(trees[f'option_{letter}'], p=p)
        else:
            table = self.doc.add_table(rows=4 // columns, cols=columns)
            table.autofit = False
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for index, letter in enumerate('abcd'):
                cell = table.cell(index // columns, index % columns)
                cell.width = Cm(17.5 / columns)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(3)
                p.add_run(f'{letter.upper()}. ').bold = True
                self.content(trees[f'option_{letter}'], cell, p)

    def exam_header(self):
        s = self.settings
        header = self.doc.add_table(rows=1, cols=2)
        header.autofit = False
        header.columns[0].width, header.columns[1].width = Cm(6), Cm(11.5)
        left, right = header.rows[0].cells
        left.width, right.width = Cm(6), Cm(11.5)
        left.text = s['organization']
        left.add_paragraph('ĐỀ THAM KHẢO')
        right.text = s['exam_title']
        right.add_paragraph('Bài thi: TOÁN')
        timing = right.add_paragraph(f'Thời gian làm bài: {s["minutes"]} phút, không kể thời gian phát đề')
        for cell in (left, right):
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs: run.bold = True
        for run in timing.runs:
            run.bold, run.italic = False, True
            run.font.size = Pt(11)
        self.paragraph(text=f'Mã đề: {s["code"]}', bold=True).alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.paragraph(text='Họ và tên thí sinh: ....................................................................................')
        self.paragraph(text='Số báo danh: ............................................................................................')

    def solution(self, question, trees):
        label = self.paragraph(text='Lời giải', bold=True, center=True)
        label.paragraph_format.keep_with_next = True
        answer = question['correct_option']
        if question['type'] == 'msq':
            answer = '; '.join(f'{letter.lower()}) {"Đúng" if letter in answer.split(",") else "Sai"}' for letter in 'ABCD')
        p = self.paragraph(text='Đáp án: ', bold=True)
        p.paragraph_format.keep_with_next = True
        if question['type'] == 'sa':
            self.content(trees['answer'], p=p)
        else:
            p.add_run(answer)
        explanation = trees['explanation']
        if ''.join(explanation.itertext()).strip() or len(explanation):
            self.content(explanation)
        else:
            self.paragraph(text='Chưa có lời giải chi tiết trong dữ liệu JSON.')

    def question_sections(self, with_solutions=False):
        questions = self.prepared['quiz']['questions']
        for kind in TYPES:
            indexes = [i for i, question in enumerate(questions) if question['type'] == kind]
            if not indexes: continue
            heading = self.paragraph(text=SECTIONS[kind], bold=True)
            heading.paragraph_format.keep_with_next = True
            suffix = {'mcq': 'Mỗi câu hỏi thí sinh chỉ chọn một phương án.', 'msq': 'Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.', 'sa': ''}[kind]
            instruction = self.paragraph(text=f'Thí sinh trả lời từ câu 1 đến câu {len(indexes)}. {suffix}'.strip())
            instruction.paragraph_format.keep_with_next = True
            for number, index in enumerate(indexes, 1):
                q, trees = questions[index], self.prepared['trees'][index]
                p = self.paragraph(text=f'Câu {number}. ', bold=True)
                p.paragraph_format.keep_with_next = True
                self.content(trees['question'], p=p)
                if kind in ('mcq', 'msq'):
                    self.options(trees, kind)
                else:
                    self.paragraph(text='Trả lời:  □ □ □ □')
                if with_solutions:
                    self.solution(q, trees)

    def build(self):
        self.paragraph(text='PHẦN 1. ĐỀ THI', bold=True, center=True).paragraph_format.keep_with_next = True
        self.exam_header()
        self.question_sections()
        self.paragraph(text='–––––––– HẾT ––––––––', bold=True, center=True)
        if self.settings['include_answers']:
            self.doc.add_page_break()
            self.paragraph(text='PHẦN 2. LỜI GIẢI CHI TIẾT', bold=True, center=True).paragraph_format.keep_with_next = True
            self.paragraph(text=f'{self.settings["exam_title"]} – Mã đề {self.settings["code"]}', bold=True, center=True)
            self.question_sections(with_solutions=True)
        output = BytesIO(); self.doc.save(output)
        return output.getvalue()


def export_quiz(source, settings, previews):
    prepared = validate_quiz(source, settings['include_answers'])
    rendered = {}
    for formula in prepared['formulas'].values():
        preview = previews.get(formula['id'])
        if not preview:
            raise ValueError('Thiếu hình xem trước công thức. Hãy kiểm tra JSON và xuất lại.')
        png, _, _ = svg_png(preview['svg'].encode(), scale=3)
        rendered[formula['id']] = (png, preview['width'], preview['height'], preview['baseline'])
    images = load_images(prepared['urls'])
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for mode, filename in [('equation', 'De-Toan-Equation.docx'), ('mathtype', 'De-Toan-MathType.docx')]:
            archive.writestr(filename, ExamWriter(prepared, settings, images, rendered, mode).build())
    if output.tell() > 4 * 1024 * 1024:
        raise ValueError('ZIP vượt quá 4 MB. Hãy giảm kích thước ảnh hoặc xuất ít câu hơn.')
    return output.getvalue()

"""End-to-end Word upload, answer analysis, multi-request ZIP and responsive UI."""
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright


def main():
    with sync_playwright() as p:
        browser = p.chromium.launch(channel='chrome', headless=True)
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto('http://127.0.0.1:5173/word-shuffle')
        page.get_by_label('Chọn đề Word').set_input_files('DE-TOAN-MAU-TRON.docx')
        page.get_by_role('button', name='Phân tích đề', exact=True).click()
        page.get_by_role('heading', name='Kết quả phân tích: 22 câu').wait_for()
        assert page.locator('.shuffle-answer', has_text='ĐSSĐ').count() == 1
        assert page.locator('.shuffle-answer', has_text='632').count() == 1
        page.get_by_label('Sở GDĐT', exact=True).fill('SỞ GDĐT HÀ NỘI')
        page.get_by_label('Trường THPT', exact=True).fill('TRƯỜNG THPT MẪU')
        page.get_by_role('button', name='Trộn đề và tạo ZIP').click()
        link = page.get_by_role('link', name='Tải ZIP đề Word và Excel đáp án')
        link.wait_for(timeout=120000)
        with page.expect_download() as download:
            link.click()
        target = Path('.artifacts/tron-de-word-browser.zip')
        download.value.save_as(target)
        with ZipFile(target) as archive:
            assert len([n for n in archive.namelist() if n.endswith('.docx')]) == 8
            from docx import Document
            teacher = Document(BytesIO(archive.read('loi_giai-0101.docx')))
            assert 'SỞ GDĐT HÀ NỘI' in teacher.tables[0].cell(0, 0).text
            assert 'TRƯỜNG THPT MẪU' in teacher.tables[0].cell(0, 0).text
            before_part_iii = '\n'.join(p.text for p in teacher.paragraphs).split('PHẦN III')[0]
            assert 'Đáp án:' not in before_part_iii
            mapping = json.loads(archive.read('Doi_chieu_cau_goc.json'))
            assert list(mapping) == ['0101', '0102', '0103', '0104']
            sheet = load_workbook(BytesIO(archive.read('Dap_an_TNMaker_2025.xlsx'))).active
            assert sheet.max_row == 23 and sheet.max_column == 5
            for column, (code, parts) in enumerate(mapping.items(), 2):
                assert sheet.cell(1, column).value == code
                assert [sheet.cell(i + 2, column).value for i in range(22)] == [q['dap_an'] for qs in parts.values() for q in qs]
        print(f'Four variants and answer workbook: {target.stat().st_size} bytes', flush=True)
        page.screenshot(path='.artifacts/word-shuffle-desktop.png', full_page=True)
        page.get_by_label('Số đề cần tạo').fill('0')
        assert page.get_by_role('button', name='Trộn đề và tạo ZIP').is_disabled()
        assert link.count() == 0
        page.get_by_label('Số đề cần tạo').fill('2')
        page.get_by_label('Trường THPT', exact=True).fill('')
        assert page.get_by_role('button', name='Trộn đề và tạo ZIP').is_disabled()
        page.get_by_label('Trường THPT', exact=True).fill('TRƯỜNG THPT MẪU')
        page.set_viewport_size({'width': 390, 'height': 844})
        page.screenshot(path='.artifacts/word-shuffle-mobile.png', full_page=True)
        assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
        page.get_by_label('Chọn đề Word').set_input_files({'name': 'invalid.docx', 'mimeType': 'application/octet-stream', 'buffer': b'broken'})
        page.get_by_role('button', name='Phân tích đề', exact=True).click()
        page.get_by_role('alert').wait_for()
        assert page.get_by_role('button', name='Trộn đề và tạo ZIP').count() == 0
        assert not errors, errors
        browser.close()


if __name__ == '__main__':
    main()

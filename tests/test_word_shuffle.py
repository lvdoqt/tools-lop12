from io import BytesIO
import json
from pathlib import Path
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from fastapi.testclient import TestClient
from lxml import etree
from openpyxl import load_workbook

from backend.main import app
from backend.services.word_shuffle import Exam, NS, export_shuffle, tag, shuffled_blocks, Question, plain


def fixture():
    doc = Document()
    doc.add_paragraph('ĐỀ TOÁN')
    for part in ('I', 'II', 'III'):
        doc.add_paragraph(f'PHẦN {part}. Kiểm tra')
        for number in range(1, 3):
            p = doc.add_paragraph('Câu ')
            p.add_run(str(number))
            p.add_run(f': Nội dung {part}-{number}')
            if part != 'III':
                labels = 'ABCD' if part == 'I' else 'abcd'
                p = doc.add_paragraph()
                for i, label in enumerate(labels):
                    run = p.add_run(label)
                    run.underline = i == number - 1
                    p.add_run(('.' if part == 'I' else ')') + f' lựa chọn {i + 1}\t')
            doc.add_paragraph('Lời giải')
            if part == 'III':
                doc.add_paragraph('Đáp án: -1.25' if number == 1 else 'Trả lời: 632')
            doc.add_paragraph(f'Giải thích {part}-{number}')
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def variant_fixture(tf_lines=None, short_text='Đáp số: 446.', punctuation='.', keep_tf_underlines=False):
    doc = Document(BytesIO(fixture()))
    part = None
    for p in doc.paragraphs:
        if p.text.startswith('PHẦN '):
            part = p.text.split()[1].rstrip('.')
        if p.text.startswith('Câu '):
            label, content = p.text.split(':', 1)
            p.clear()
            for token in (label[:4], label[4:], punctuation):
                p.add_run(token).bold = True
            p.add_run(content)
        if part == 'II':
            if not keep_tf_underlines:
                for run in p.runs:
                    run.underline = False
            if p.text == 'Lời giải' and tf_lines is not None:
                for line in reversed(tf_lines):
                    following = doc.add_paragraph()
                    # Split verdict labels and text across bold Word runs.
                    following.add_run(line[:2]).bold = True
                    following.add_run(line[2:]).bold = True
                    p._p.addnext(following._p)
        if part == 'III' and p.text.startswith(('Đáp án:', 'Trả lời:')):
            p.clear()
            p.add_run(short_text).bold = True
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


class WordShuffleTests(unittest.TestCase):
    def test_bold_question_labels_accept_dot_and_colon(self):
        for punctuation in ('.', ':'):
            with self.subTest(punctuation=punctuation):
                content = variant_fixture(punctuation=punctuation)
                exam = Exam(content)
                self.assertEqual(exam.errors, [])
                self.assertEqual([len(q) for q in exam.sections.values()], [2, 2, 2])
                archive = ZipFile(BytesIO(export_shuffle(content, 1, seed=3)))
                output = Exam(archive.read('loi_giai-0101.docx'))
                self.assertEqual(output.errors, [])
                for questions in output.sections.values():
                    for q in questions:
                        self.assertTrue(plain(q.blocks).startswith(f'Câu {q.number}{punctuation}'))

    def test_short_answer_aliases_and_terminal_period(self):
        for label in ('Đáp án', 'Đáp số', 'Trả lời'):
            for number, expected in (('446.', '446'), ('446', '446'), ('5.19.', '5,19'),
                                     ('5,19.', '5,19'), ('-0.25.', '-0,25'), ('0.', '0')):
                with self.subTest(label=label, number=number):
                    exam = Exam(variant_fixture(short_text=f'{label}: {number}'))
                    self.assertEqual(exam.errors, [])
                    self.assertEqual([q.answer for q in exam.sections['III']], [expected, expected])

    def test_tf_solution_verdicts_preview_teacher_and_excel(self):
        for lines in (['a) ĐÚNG', 'b) SAI.', 'c) Đúng. Vì kết quả đúng.', 'd) sai'],
                      ['a) ĐÚNG b) SAI c) Đúng d) Sai']):
            with self.subTest(lines=lines):
                content = variant_fixture(tf_lines=lines)
                exam = Exam(content)
                self.assertEqual(exam.errors, [])
                self.assertEqual([q.answer for q in exam.sections['II']], ['ĐSĐS', 'ĐSĐS'])
                self.assertEqual([o['correct'] for o in exam.sections['II'][0].options], [True, False, True, False])
                archive = ZipFile(BytesIO(export_shuffle(content, 1, seed=9)))
                teacher = Exam(archive.read('loi_giai-0101.docx'))
                self.assertEqual(teacher.errors, [])
                self.assertEqual([q.answer for q in teacher.sections['II']], ['ĐSĐS', 'ĐSĐS'])
                from backend.services.word_shuffle import text_map
                for q in teacher.sections['II']:
                    p = next(p for p in q.blocks if text_map(p)[0].startswith('a)'))
                    value, owners = text_map(p)
                    for label, correct in zip('abcd', (True, False, True, False)):
                        self.assertEqual(teacher.underlined(owners[value.index(label + ')')]), correct)
                sheet = load_workbook(BytesIO(archive.read('Dap_an_TNMaker_2025.xlsx'))).active
                self.assertEqual([sheet.cell(row, 2).value for row in (4, 5, 6, 7)], ['ĐSĐS', 'ĐSĐS', '446', '446'])
                student = Document(BytesIO(archive.read('De_0101.docx')))
                self.assertFalse(any(r.underline for p in student.paragraphs for r in p.runs))

    def test_tf_missing_or_contradictory_solution_is_not_silently_scored(self):
        for lines, marks, message in ((['a) ĐÚNG'], False, 'chưa ghi'),
                                      (['a) ĐÚNG b) SAI c) SAI d) ĐÚNG', 'a) SAI'], False, 'cả ĐÚNG và SAI'),
                                      (['a) SAI b) SAI c) SAI d) SAI'], True, 'gạch chân')):
            with self.subTest(lines=lines):
                content = variant_fixture(tf_lines=lines, keep_tf_underlines=marks)
                self.assertTrue(any(message in e for e in Exam(content).errors))
                with self.assertRaises(ValueError):
                    export_shuffle(content, 1)

    def test_answers_across_runs_and_numeric_normalization(self):
        exam = Exam(fixture())
        self.assertEqual(exam.errors, [])
        self.assertEqual([q.answer for q in exam.sections['I']], ['A', 'B'])
        self.assertEqual([q.answer for q in exam.sections['II']], ['ĐSSS', 'SĐSS'])
        self.assertEqual([q.answer for q in exam.sections['III']], ['-1,25', '632'])

    def test_unique_variants_sections_numbering_and_excel_alignment(self):
        content = fixture()
        archive = ZipFile(BytesIO(export_shuffle(content, 8, '0091', seed=17)))
        mapping = json.loads(archive.read('Doi_chieu_cau_goc.json'))
        workbook = load_workbook(BytesIO(archive.read('Dap_an_TNMaker_2025.xlsx')))
        self.assertEqual(workbook.sheetnames, ['Dữ liệu'])
        sheet = workbook.active
        self.assertEqual(sheet['A1'].value, 'Câu\\Mã đề')
        signatures = set()
        for column, (code, parts) in enumerate(mapping.items(), 2):
            self.assertEqual(sheet.cell(1, column).value, code)
            answers = [q['dap_an'] for part in parts.values() for q in part]
            self.assertEqual([sheet.cell(i + 2, column).value for i in range(6)], answers)
            signatures.add(tuple((q['cau_goc'], tuple(q.get('phuong_an_moi_sang_goc', {}).values())) for part in parts.values() for q in part))
            teacher = Exam(archive.read(f'loi_giai-{code}.docx'))
            self.assertEqual(teacher.errors, [])
            for part, questions in teacher.sections.items():
                self.assertEqual([q.number for q in questions], [1, 2])
                self.assertEqual([q.answer for q in questions], [q['dap_an'] for q in parts[part]])
                for q, original in zip(questions, parts[part]):
                    self.assertIn(f'Nội dung {part}-{original["cau_goc"]}', ''.join(q.blocks[0].itertext()))
                    if part == 'I':
                        for option in q.options:
                            old_label = original['phuong_an_moi_sang_goc'][option['label']]
                            self.assertEqual(option['text'], f'lựa chọn {"ABCD".index(old_label) + 1}')
                        self.assertEqual(original['phuong_an_moi_sang_goc'][q.answer], 'ABCD'[original['cau_goc'] - 1])
                    if part in ('I', 'II'):
                        self.assertNotIn('Đáp án:', plain(q.blocks + q.solution))
            student = Document(BytesIO(archive.read(f'De_{code}.docx')))
            all_text = '\n'.join(p.text for p in student.paragraphs)
            self.assertNotIn('Đáp án:', all_text)
            self.assertNotIn('Trả lời:', all_text)
            self.assertNotIn('Lời giải', all_text)
            self.assertFalse(any(r.underline for p in student.paragraphs for r in p.runs))
        self.assertEqual(len(signatures), 8)

    def test_more_variants_than_question_orders(self):
        archive = ZipFile(BytesIO(export_shuffle(fixture(), 10, seed=19)))
        mapping = json.loads(archive.read('Doi_chieu_cau_goc.json'))
        self.assertEqual(len(mapping), 10)

    def test_official_heading_and_red_code(self):
        archive = ZipFile(BytesIO(export_shuffle(fixture(), 1, '0000', seed=1,
                                                department='Sở GDĐT Hà Nội', school='Trường THPT Mẫu')))
        for name in ('De_0000.docx', 'loi_giai-0000.docx'):
            doc = Document(BytesIO(archive.read(name)))
            self.assertEqual(doc.tables[0].cell(0, 0).text, 'SỞ GDĐT HÀ NỘI\nTRƯỜNG THPT MẪU')
            self.assertEqual(doc.tables[0].cell(0, 1).text, 'ĐỀ THI THỬ TN THPT 2027\nMÔN: TOÁN\nThời gian làm bài: 90 phút')
            code = next(p for p in doc.paragraphs if p.text == 'Mã đề: 0000')
            self.assertEqual(code.runs[0].font.size.pt, 18)
            self.assertEqual(str(code.runs[0].font.color.rgb), 'FF0000')
            self.assertNotIn('ĐỀ TOÁN', '\n'.join(p.text for p in doc.paragraphs))

    def test_multiline_options_and_explicit_solution_reference(self):
        doc = Document()
        doc.add_paragraph('Câu 1: Giữ câu hỏi')
        for i, label in enumerate('ABCD'):
            p = doc.add_paragraph()
            p.add_run(label).underline = label == 'B'
            p.add_run(f'. lựa chọn {i}')
            doc.add_paragraph(f'dòng tiếp {label}')
        q = Question(1, blocks=[p._p for p in doc.paragraphs], option_order=(3, 2, 1, 0))
        blocks = shuffled_blocks(q)
        self.assertEqual(plain(blocks), 'Câu 1: Giữ câu hỏi\nA. lựa chọn 3\ndòng tiếp D\nB. lựa chọn 2\ndòng tiếp C\nC. lựa chọn 1\ndòng tiếp B\nD. lựa chọn 0\ndòng tiếp A')
        from backend.services.word_shuffle import remap_solution
        p = doc.add_paragraph('Chọn ')
        p.add_run('B')
        p.add_run('. Điểm B thuộc đường thẳng AB.')
        self.assertEqual(plain(remap_solution([p._p], q.option_order)), 'Chọn C. Điểm B thuộc đường thẳng AB.')

    def test_batched_exports_have_same_plan_and_workbook(self):
        content = fixture()
        full = ZipFile(BytesIO(export_shuffle(content, 4, seed='session')))
        for index in range(4):
            batch = ZipFile(BytesIO(export_shuffle(content, 4, seed='session', variant_index=index)))
            self.assertEqual(json.loads(batch.read('Doi_chieu_cau_goc.json')), json.loads(full.read('Doi_chieu_cau_goc.json')))
            self.assertEqual(len([n for n in batch.namelist() if n.endswith('.docx')]), 2)
            for name in batch.namelist():
                if name.endswith('.docx'):
                    self.assertEqual(ZipFile(BytesIO(batch.read(name))).read('word/document.xml'),
                                     ZipFile(BytesIO(full.read(name))).read('word/document.xml'))
            self.assertEqual(list(load_workbook(BytesIO(batch.read('Dap_an_TNMaker_2025.xlsx'))).active.values),
                             list(load_workbook(BytesIO(full.read('Dap_an_TNMaker_2025.xlsx'))).active.values))

    def test_invalid_answers_and_files_block_export(self):
        doc = Document(BytesIO(fixture()))
        for p in doc.paragraphs:
            for r in p.runs:
                r.underline = False
        out = BytesIO(); doc.save(out)
        self.assertTrue(Exam(out.getvalue()).errors)
        with self.assertRaisesRegex(ValueError, 'gạch chân'):
            export_shuffle(out.getvalue())
        with self.assertRaises(ValueError):
            Exam(b'not a docx')
        with self.assertRaises(ValueError):
            export_shuffle(fixture(), 3, '9999')
        with self.assertRaises(ValueError):
            export_shuffle(fixture(), 3, seed=None, variant_index=1)

    def test_zip_expansion_guard(self):
        out = BytesIO()
        with ZipFile(out, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('huge', b'0' * (65 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, 'sau giải nén'):
            Exam(out.getvalue())

    def test_api_upload_validation_and_export(self):
        client = TestClient(app)
        content = fixture()
        files = {'file': ('exam.docx', content)}
        response = client.post('/api/word-shuffle/analyze', files=files)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['total'], 6)
        response = client.post('/api/word-shuffle/export', files=files,
                               data={'count': 2, 'start_code': '001', 'include_solutions': 'false', 'seed': 'test', 'variant_index': 1})
        self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
        archive = ZipFile(BytesIO(response.content))
        self.assertIn('De_002.docx', archive.namelist())
        self.assertFalse(any(n.startswith('loi_giai-') for n in archive.namelist()))
        self.assertEqual(client.post('/api/word-shuffle/analyze', files={'file': ('bad.doc', content)}).status_code, 400)
        self.assertEqual(client.post('/api/word-shuffle/analyze', files={'file': ('bad.docx', b'')}).status_code, 400)
        self.assertEqual(client.post('/api/word-shuffle/export', files=files, data={'count': 0}).status_code, 422)
        self.assertEqual(client.post('/api/word-shuffle/export', files=files, data={'variant_index': 4, 'count': 2, 'seed': 'x'}).status_code, 400)

    @unittest.skipUnless(Path('DE-TOAN-MAU-TRON.docx').exists(), 'Local user sample')
    def test_user_sample_answers_and_preserved_embedded_resources(self):
        content = Path('DE-TOAN-MAU-TRON.docx').read_bytes()
        exam = Exam(content)
        self.assertEqual(exam.errors, [])
        self.assertEqual([len(q) for q in exam.sections.values()], [12, 4, 6])
        self.assertEqual([q.answer for q in exam.sections['I']], list('CBABABABACDD'))
        self.assertEqual([q.answer for q in exam.sections['II']], ['ĐSSĐ', 'ĐSĐS', 'ĐĐSĐ', 'SĐSĐ'])
        self.assertEqual([q.answer for q in exam.sections['III']], ['3,48', '632', '147', '126', '5,19', '1,92'])
        original = ZipFile(BytesIO(content))
        archive = ZipFile(BytesIO(export_shuffle(content, 4, seed=7, variant_index=0)))
        teacher = ZipFile(BytesIO(archive.read('loi_giai-0101.docx')))
        student = ZipFile(BytesIO(archive.read('De_0101.docx')))
        for name in original.namelist():
            if name != 'word/document.xml':
                self.assertEqual(original.read(name), teacher.read(name))
                self.assertEqual(original.read(name), student.read(name))
        for name in ('word/document.xml',):
            old = etree.fromstring(original.read(name))
            new = etree.fromstring(teacher.read(name))
            self.assertEqual(len(old.xpath('.//w:object', namespaces=NS)), len(new.xpath('.//w:object', namespaces=NS)))
        root = etree.fromstring(student.read('word/document.xml'))
        self.assertFalse(root.findall('.//w:highlight', NS))
        self.assertTrue(all(u.get(tag('val')) == 'none' for u in root.findall('.//w:u', NS)))
        teacher_exam = Exam(archive.read('loi_giai-0101.docx'))
        self.assertEqual(teacher_exam.errors, [])
        mapping = json.loads(archive.read('Doi_chieu_cau_goc.json'))['0101']
        # Match each moved answer's exact embedded equation IDs against its old label.
        def option_objects(question):
            from backend.services.word_shuffle import MCQ_LABEL, text_map, rich_slice, paragraphs
            result = {}
            for p in paragraphs(question.blocks):
                value = text_map(p)[0]
                matches = list(MCQ_LABEL.finditer(value))
                for i, match in enumerate(matches):
                    fragment = rich_slice(p, match.start(), matches[i + 1].start() if i + 1 < len(matches) else len(value))
                    result[match[1]] = fragment.xpath('.//@r:id', namespaces={'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})
            return result
        for part in ('I', 'II', 'III'):
            for question, item in zip(teacher_exam.sections[part], mapping[part]):
                old = next(q for q in exam.sections[part] if q.number == item['cau_goc'])
                self.assertEqual(question.answer, item['dap_an'])
                if part == 'I':
                    source_objects, new_objects = option_objects(old), option_objects(question)
                    for new_label, old_label in item['phuong_an_moi_sang_goc'].items():
                        self.assertEqual(new_objects[new_label], source_objects[old_label])
                    self.assertEqual(item['phuong_an_moi_sang_goc'][question.answer], old.answer)
                else:
                    self.assertEqual(question.answer, old.answer)
                    self.assertEqual(plain(question.solution), plain(old.solution))


if __name__ == '__main__':
    unittest.main()
